#!/usr/bin/env python3
"""Sinh GreenSM-Pitch-10min.pptx — 12 slide / 10 phút. Chạy:
uv run --with python-pptx python scripts/slides/build_greensm_pitch.py"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

GREEN = RGBColor(0x0E, 0x7C, 0x55)
INK   = RGBColor(0x17, 0x1D, 0x1A)
MUTED = RGBColor(0x5B, 0x65, 0x5F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
A = "docs/project/slides/assets/pitch"

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def txt(slide, x, y, w, h, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.alignment = align
        for r in p.runs: r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def kicker(slide, label):          # nhãn phần nhỏ phía trên tiêu đề
    txt(slide, 0.6, 0.35, 8, 0.4, label, 13, GREEN, bold=True)

def title(slide, text, size=32):
    txt(slide, 0.6, 0.75, 12.1, 1.0, text, size, INK, bold=True)

def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s

def pic_fit(slide, path, x, y, w):  # ảnh giữ tỉ lệ theo bề ngang
    assert os.path.exists(path), f"missing asset: {path} (chạy từ repo root?)"
    slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))

def stat(slide, x, y, big, small):
    txt(slide, x, y, 3.0, 1.0, big, 44, GREEN, bold=True)
    txt(slide, x, y+0.85, 3.0, 0.6, small, 14, MUTED)

def bullets(slide, items, x=0.7, y=1.9, w=12.0, size=20, gap=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5.2))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        r = p.add_run(); r.text = "▪ " + head + "  "; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = INK
        r2 = p.add_run(); r2.text = body; r2.font.size = Pt(size - 2); r2.font.color.rgb = MUTED

# ---- S1 HOOK ----
s = prs.slides.add_slide(BLANK)
pic_fit(s, f"{A}/greensm-photo.jpg", 6.9, 0, 6.45)          # ảnh phải, full-bleed dọc
txt(s, 0.6, 0.5, 6.0, 0.5, "AI20K · COHORT 2 · TEAM 116", 13, GREEN, bold=True)
txt(s, 0.6, 1.3, 6.1, 2.2, "2 GIỜ SÁNG.\nAI ĐANG ĐỌC\nREVIEW CỦA BẠN?", 40, INK, bold=True)
stat(s, 0.6, 4.0, "4.8★", "App Store Việt Nam")
stat(s, 3.4, 4.0, "1,5 TRIỆU", "lượt đánh giá — ai đọc hết?")
txt(s, 0.6, 5.6, 6.1, 1.2, "AI Quality Intelligence Platform cho Green SM\nBiến 1,5 triệu tiếng nói thành hành động — tự động, có kiểm soát.", 16, MUTED)
notes(s, "0:00–0:40 · Im lặng 2s. Câu chuyện 2 giờ sáng → con số 4.8★/1,5 triệu → 'ai đọc hết?' → giới thiệu team 116.")

# ---- S2 BẰNG CHỨNG THẬT ----
s = prs.slides.add_slide(BLANK)
kicker(s, "PHẦN 1 · BÀI TOÁN"); title(s, "Điểm cao không có nghĩa là an toàn")
for i, (q, tag) in enumerate([
    ("“Tài xế có hành vi hành hung hành khách…”", "Review thật · App Store"),
    ("“Không đăng nhập được app trên máy của tôi.”", "Review thật · App Store"),
    ("“Tài xế say xỉn gây tai nạn, xe tông người rồi bỏ chạy.”", "Review thật · dữ liệu dự án"),
]):
    box = s.shapes.add_textbox(Inches(0.7), Inches(1.9 + i*1.15), Inches(7.3), Inches(1.0))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = q; r.font.size = Pt(18); r.font.italic = True; r.font.color.rgb = INK
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = tag; r2.font.size = Pt(12); r2.font.color.rgb = GREEN
stat(s, 8.6, 2.0, "34", "tỉnh thành phủ sóng")
stat(s, 10.9, 2.0, "4", "quốc gia")
stat(s, 8.6, 3.7, "1000+", "review mới mỗi ngày")
txt(s, 8.6, 5.3, 4.2, 1.2, "Ca nghiêm trọng chìm giữa\nđám đông hài lòng.", 16, INK, bold=True)
notes(s, "0:40–1:40 · Đọc 2 review thật trên App Store, nhấn: chúng nằm ngay dưới điểm 4.8★. Quy mô 34 tỉnh/4 quốc gia → rủi ro ở phần đuôi.")

# ---- S3 BA LỖ HỔNG ----
s = prs.slides.add_slide(BLANK)
kicker(s, "PHẦN 1 · BÀI TOÁN"); title(s, "Ba lỗ hổng khiến 2 giờ sáng đơn độc")
bullets(s, [
    ("Quá tải.", "Hàng nghìn review/ngày — không đội ngũ nào đọc xuể bằng tay."),
    ("An toàn bị bỏ sót.", "Review 'tài xế say rượu' trông y hệt review 'app lag' — cho đến khi có người đọc kỹ. Người đọc kỹ luôn đến sau."),
    ("Sai địa chỉ.", "Vấn đề rơi vào hộp thư chung: ai cũng thấy, không ai chịu trách nhiệm."),
], y=2.1, size=22)
notes(s, "1:40–2:30 · Không phải thiếu dữ liệu — dữ liệu thừa mứa. Nhấn lỗ hổng 2 là nguy hiểm nhất.")

# ---- S4 ĐIỀU CẦN TỒN TẠI ----
s = prs.slides.add_slide(BLANK)
kicker(s, "PHẦN 1 · BÀI TOÁN"); title(s, "Điều cần tồn tại — và vì sao là bây giờ")
txt(s, 0.7, 2.2, 12.0, 1.6, "Một hệ thống ĐỌC THAY con người,\nPHÂN LOẠI đúng mức khẩn cấp, GIAO ĐÚNG NGƯỜI xử lý.", 30, GREEN, bold=True)
bullets(s, [
    ("Why now:", "AI đọc hiểu tiếng Việt đã đủ tốt và đủ rẻ để đọc 1,5 triệu review — 3 năm trước còn bất khả thi."),
    ("Và nó không phải ý tưởng:", "hệ thống đã chạy thật — 5 phút tiếp theo là bằng chứng."),
], y=4.3, size=20)
notes(s, "2:30–3:00 · Chốt phần bài toán bằng 1 câu giải pháp; bridge sang phần sản phẩm.")

# ---- S5 HÀNH TRÌNH 7 CHẶNG ----
s = prs.slides.add_slide(BLANK)
kicker(s, "PHẦN 2 · SẢN PHẨM"); title(s, "Phòng cấp cứu cho review — hành trình 7 chặng")
steps = ["① Che thông tin\nnhạy cảm", "② AI đọc lướt\n(topic·cảm xúc·mức độ)", "③ Luật an toàn\nghi đè", "④ Phát hiện\nổ dịch", "⑤ Định tuyến\n6 phòng ban", "⑥ Lập việc\ntự động/duyệt", "⑦ Người ký\nxác nhận"]
for i, st_ in enumerate(steps):
    x = 0.45 + i * 1.82
    box = s.shapes.add_textbox(Inches(x), Inches(2.5), Inches(1.72), Inches(1.6))
    tf = box.text_frame; tf.word_wrap = True
    for j, line in enumerate(st_.split("\n")):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = line; r.font.size = Pt(13 if j else 15); r.font.bold = (j == 0)
        r.font.color.rgb = GREEN if j == 0 else MUTED
bullets(s, [
    ("Rẻ chạy trước:", "ML nhỏ xử lý mọi review; LLM chỉ gọi khi cần suy luận sâu."),
    ("An toàn thắng xác suất.", "Luật an toàn luôn ghi đè được AI."),
    ("Đổi luật = sửa 1 file cấu hình.", "Một khối hỏng, hệ thống vẫn đứng."),
], y=4.7, size=17)
notes(s, "3:00–3:50 · Giới thiệu ẩn dụ phòng cấp cứu; lướt 7 chặng; 4 nguyên tắc vận hành.")

# ---- S6..S9: split ảnh phải ----
def split_slide(kick, ttl, items, img, note):
    s = prs.slides.add_slide(BLANK)
    kicker(s, kick); title(s, ttl, size=28)
    bullets(s, items, x=0.7, y=1.9, w=6.6, size=17)
    pic_fit(s, img, 7.6, 1.7, 5.4)
    notes(s, note); return s

split_slide("PHẦN 2 · SẢN PHẨM", "Chặng 1–3 · Làm sạch → AI đọc lướt → Luật ghi đè", [
    ("Che PII trước tiên.", "SĐT, email, biển số — trước khi người hay máy đọc."),
    ("AI nhanh & rẻ, <0,1s/review.", "77,1% chủ đề · 94,3% cảm xúc · hiểu phủ định tiếng Việt."),
    ("Ví dụ thật:", "“Tài xế say xỉn gây tai nạn… bỏ chạy” — AI đoán mức trung bình, luật an toàn GHI ĐÈ → khẩn cấp cao nhất, bắt buộc người xem."),
    ("Nguyên tắc:", "khi nghi ngờ an toàn, máy không được tự tin hơn sự thận trọng."),
], f"{A}/ui-feedback.png", "3:50–4:50 · Đi chậm ở ví dụ thật; chỉ vào screenshot màn Phản hồi.")

split_slide("PHẦN 2 · SẢN PHẨM", "Chặng 4–5 · Ổ dịch → Báo động → Đúng chuyên khoa", [
    ("Ngưỡng rõ ràng:", "1 ca khẩn cấp nhất → báo NGAY · ≥3 ca mức cao cùng khu vực → cảnh báo · ≥10 ca tiêu cực → báo động đỏ · chờ 60 phút mới báo lại."),
    ("Bảng phân khoa 6 phòng ban, kèm SLA.", "An toàn → vận hành trong 1 GIỜ; lỗi app → kỹ thuật 12h; thanh toán → CSKH 12h."),
    ("Đổi luật phân khoa = sửa 1 file cấu hình.", "Không viết lại dòng code nào, không triển khai lại."),
], f"{A}/ui-ticket.png", "4:50–5:40 · Chỉ màn Phân loại ticket; nhấn câu 'không viết lại code'.")

split_slide("PHẦN 2 · SẢN PHẨM", "Chặng 6 · Từ báo động đến việc thật", [
    ("Đường tự động:", "gom báo động cùng phòng ban + khu vực → tự lập đầu việc; mỗi cụm chỉ 1 việc đang mở — không trùng, không rơi rớt."),
    ("Đường có kiểm soát:", "AI viết báo cáo đề xuất (nguyên nhân gốc + việc nháp) → quản lý thêm/sửa/xóa → bấm 1 nút duyệt = duyệt + tạo việc trong cùng một hành động, kèm link ngược báo cáo gốc."),
    ("MÁY ĐỀ XUẤT. NGƯỜI QUYẾT ĐỊNH.", "Hệ thống thực thi trọn vẹn."),
], f"{A}/ui-report.png", "5:40–6:30 · Màn Báo cáo AI; câu đinh: máy đề xuất, người quyết định.")

split_slide("PHẦN 2 · SẢN PHẨM", "Chặng 7 · Con người trong vòng lặp", [
    ("4 tình huống bắt buộc người xem:", "khẩn cấp nhất / từ khóa an toàn · AI không chắc · câu mơ hồ · có người yêu cầu."),
    ("Người duyệt: đồng ý · sửa nhãn · đánh dấu xong.", "Mỗi lần sửa ghi sổ: ai, trường nào, nhãn cũ → mới, lúc nào."),
    ("Cuốn sổ 2 vai trò:", "hồ sơ kiểm tra lại + dữ liệu cho AI học lại."),
    ("VÒNG HỌC KHÉP KÍN:", "AI gán sai → người sửa → hệ thống ghi nhớ → AI giỏi hơn."),
], f"{A}/ui-feedback.png", "6:30–7:20 · Chỉ vào hàng đợi 'Cần duyệt' + nút 'sửa nhãn' trên screenshot.")

# ---- S10 BẰNG CHỨNG SỐ ----
s = prs.slides.add_slide(BLANK)
kicker(s, "PHẦN 2 · SẢN PHẨM"); title(s, "Chạy thật, đo thật — không phải số trình diễn")
for i, (big, small) in enumerate([
    ("77,1%", "đúng chủ đề\n(test 932 mẫu)"), ("94,3%", "đúng cảm xúc"), ("97%", "không bỏ sót\nphản hồi tiêu cực"),
    ("97 giây", "xử lý 4.837 review"), ("6/6", "kịch bản kiểm thử đạt"),
]):
    stat(s, 0.7 + (i % 3) * 4.2, 2.2 + (i // 3) * 2.3, big, small)
txt(s, 0.7, 6.5, 12.0, 0.6, "Mọi con số tái lập được bằng bộ kiểm thử trong mã nguồn.", 16, INK, bold=True)
notes(s, "7:20–8:00 · Đọc chậm từng số, câu đời thường trước – phần trăm sau.")

# ---- S11 HẠN CHẾ & ROADMAP ----
s = prs.slides.add_slide(BLANK)
kicker(s, "PHẦN 3 · KẾT"); title(s, "Chúng tôi biết mình đang ở đâu")
bullets(s, [
    ("Khóa sâu tầng dữ liệu.", "Cửa chính (tầng ứng dụng) đã gate chặt; lớp khóa tại kho dữ liệu là bước kế — biết chính xác việc cần làm."),
    ("Nối dữ liệu sống.", "Đang chạy tập cố định ~5.000 review thật; kế tiếp: ống dữ liệu hằng ngày, giữ nguyên phần xử lý."),
    ("Báo động mới chỉ trên web.", "Người trực không mở máy 2 giờ sáng vẫn lỡ. Kế tiếp: đẩy thẳng vào điện thoại."),
], y=2.1, size=20)
notes(s, "8:00–8:50 · Giọng tự tin, không xin lỗi. Ý 3 móc lại câu chuyện mở đầu — nhấn chậm.")

# ---- S12 CLOSING / VISION ----
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)  # 1 = RECTANGLE
bg.fill.solid(); bg.fill.fore_color.rgb = GREEN; bg.line.fill.background()
txt(s, 0.9, 1.1, 11.5, 1.6, "Vòng lặp đã khép.", 40, WHITE, bold=True)
txt(s, 0.9, 2.5, 11.5, 0.7, "📥 Review vào  →  🤖 AI phân loại  →  🔔 Báo động  →  ✅ Quản lý duyệt  →  💼 Nhân viên xử lý  →  📋 Hoàn thành  ↺", 15, WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, 0.9, 3.6, 11.5, 1.3, "Một lời phàn nàn không được đọc là một sự cố đang chờ xảy ra lần nữa.\nTầm nhìn 3–5 năm: tầng lắng nghe tiêu chuẩn cho mọi nền tảng tiêu dùng Việt Nam —\nnơi mọi lời phàn nàn được đọc, được định tuyến, và trở thành hành động.", 18, WHITE)
txt(s, 0.9, 5.6, 11.5, 1.0, "Với chúng tôi — nó sẽ được đọc.   ·   Team 116 · AI20K Cohort 2 · Q&A", 20, WHITE, bold=True)
notes(s, "8:50–10:00 · Callback 2 giờ sáng → đọc vòng lặp 6 bước → tầm nhìn → câu kết → mời Q&A.")

out = "docs/project/slides/GreenSM-Pitch-10min.pptx"
prs.save(out); print("saved", out, "slides:", len(prs.slides._sldIdLst))
