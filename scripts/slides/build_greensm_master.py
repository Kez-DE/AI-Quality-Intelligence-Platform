#!/usr/bin/env python3
"""Sinh GreenSM-Master-Deck.pptx — 25 slide / ~25 phút. Chạy từ repo root:
uv run --with python-pptx python scripts/slides/build_greensm_master.py"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

GREEN      = RGBColor(0x0E, 0x7C, 0x55)
GREEN_SOFT = RGBColor(0xE3, 0xF1, 0xEA)
INK        = RGBColor(0x17, 0x1D, 0x1A)
MUTED      = RGBColor(0x5B, 0x65, 0x5F)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
AMBER      = RGBColor(0xB4, 0x53, 0x09)
AMBER_SOFT = RGBColor(0xFE, 0xF3, 0xC7)
ROW_ALT    = RGBColor(0xF4, 0xF7, 0xF5)
A = "docs/project/slides/assets/pitch"

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
PAGE = [0]  # đếm trang cho footer

def txt(slide, x, y, w, h, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.italic = italic
    return tb

def rect(slide, x, y, w, h, fill, line_color=None, rounded=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded: shp.adjustments[0] = 0.08
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_color is None: shp.line.fill.background()
    else: shp.line.color.rgb = line_color; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp

def new_slide(kick=None, ttl=None, ttl_size=30):
    s = prs.slides.add_slide(BLANK); PAGE[0] += 1
    rect(s, 0, 0, 13.333, 0.12, GREEN)                      # band trên
    if kick: txt(s, 0.6, 0.32, 9, 0.4, kick, 12, GREEN, bold=True)
    if ttl:  txt(s, 0.6, 0.68, 12.1, 1.0, ttl, ttl_size, INK, bold=True)
    txt(s, 0.6, 7.08, 10.0, 0.35, "Green SM · AI Quality Intelligence Platform · Team 116", 10, MUTED)
    txt(s, 12.2, 7.08, 0.9, 0.35, str(PAGE[0]), 10, MUTED, align=PP_ALIGN.RIGHT)
    return s

def notes(slide, txt_): slide.notes_slide.notes_text_frame.text = txt_

def pic_fit(slide, path, x, y, w):
    assert os.path.exists(path), f"missing asset: {path} (chạy từ repo root?)"
    slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))

def card(slide, x, y, w, h, head, body, fill=GREEN_SOFT, head_color=GREEN, body_size=13):
    rect(slide, x, y, w, h, fill, line_color=head_color, rounded=True)
    tb = slide.shapes.add_textbox(Inches(x+0.15), Inches(y+0.1), Inches(w-0.3), Inches(h-0.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = head
    r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = head_color
    for line in body.split("\n"):
        p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = line
        r2.font.size = Pt(body_size); r2.font.color.rgb = INK

def stat_card(slide, x, y, w, big, small):
    rect(slide, x, y, w, 1.55, GREEN_SOFT, line_color=GREEN, rounded=True)
    txt(slide, x+0.12, y+0.12, w-0.24, 0.8, big, 34, GREEN, bold=True)
    txt(slide, x+0.12, y+0.95, w-0.24, 0.55, small, 12, MUTED)

def bullets(slide, items, x=0.7, y=1.9, w=12.0, size=18):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(5.0))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        r = p.add_run(); r.text = "▪ " + head + "  "; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = INK
        r2 = p.add_run(); r2.text = body; r2.font.size = Pt(size-2); r2.font.color.rgb = MUTED

def table(slide, x, y, w, h, headers, rows, col_w=None, font=12):
    shp = slide.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tb = shp.table
    if col_w:
        for i, cw in enumerate(col_w): tb.columns[i].width = Inches(cw)
    for j, htxt in enumerate(headers):
        c = tb.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = GREEN
        c.text_frame.text = htxt
        for r in c.text_frame.paragraphs[0].runs:
            r.font.size = Pt(font); r.font.bold = True; r.font.color.rgb = WHITE
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tb.cell(i, j); c.fill.solid()
            c.fill.fore_color.rgb = ROW_ALT if i % 2 == 0 else WHITE
            c.text_frame.text = val; c.text_frame.word_wrap = True
            for r in c.text_frame.paragraphs[0].runs:
                r.font.size = Pt(font); r.font.color.rgb = INK

def chips(slide, steps, y=2.35, x0=0.45, cw=1.72, gap=0.1):
    for i, st in enumerate(steps):
        x = x0 + i * (cw + gap)
        rect(slide, x, y, cw, 1.5, GREEN_SOFT, line_color=GREEN, rounded=True)
        tb = slide.shapes.add_textbox(Inches(x+0.06), Inches(y+0.08), Inches(cw-0.12), Inches(1.35))
        tf = tb.text_frame; tf.word_wrap = True
        for j, line in enumerate(st.split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = line
            r.font.size = Pt(14 if j == 0 else 11); r.font.bold = (j == 0)
            r.font.color.rgb = GREEN if j == 0 else MUTED

# ================= HỒI 1 · CÂU CHUYỆN =================
# S1 HOOK (không dùng new_slide — full design riêng)
s = prs.slides.add_slide(BLANK); PAGE[0] += 1
pic_fit(s, f"{A}/greensm-photo.jpg", 6.9, 0, 6.45)
rect(s, 0, 7.1, 6.9, 0.4, GREEN)
txt(s, 0.6, 0.5, 6.0, 0.5, "AI20K · COHORT 2 · TEAM 116 · MASTER DECK", 13, GREEN, bold=True)
txt(s, 0.6, 1.25, 6.1, 2.2, "2 GIỜ SÁNG.\nAI ĐANG ĐỌC\nREVIEW CỦA BẠN?", 40, INK, bold=True)
stat_card(s, 0.6, 3.95, 2.7, "4.8★", "App Store Việt Nam")
stat_card(s, 3.5, 3.95, 3.1, "1,5 TRIỆU", "lượt đánh giá — ai đọc hết?")
txt(s, 0.6, 5.8, 6.1, 1.1, "AI Quality Intelligence Platform cho Green SM\nBiến 1,5 triệu tiếng nói thành hành động — tự động, có kiểm soát.", 15, MUTED)
notes(s, "~60s · Im lặng 2s. Chuyện 2 giờ sáng → 4.8★/1,5 triệu → 'ai đọc hết?' → hứa hẹn 25 phút: từ nỗi đau thật đến hệ thống chạy thật.")

# S2 TĂNG TRƯỞNG + REVIEW THẬT
s = new_slide("HỒI 1 · CÂU CHUYỆN", "Tăng trưởng nóng — và 1,5 triệu tiếng nói không ai đọc")
for i, (q, tag) in enumerate([
    ("“Tài xế có hành vi hành hung hành khách…”", "Review thật · App Store"),
    ("“Không đăng nhập được app trên máy của tôi.”", "Review thật · App Store"),
    ("“Tài xế say xỉn gây tai nạn, xe tông người rồi bỏ chạy.”", "Review thật · dữ liệu dự án"),
]):
    rect(s, 0.7, 1.95 + i*1.5, 7.2, 1.3, AMBER_SOFT, line_color=AMBER, rounded=True)
    txt(s, 0.9, 2.08 + i*1.5, 6.8, 0.7, q, 15, INK, italic=True)
    txt(s, 0.9, 2.72 + i*1.5, 6.8, 0.4, tag, 11, AMBER, bold=True)
stat_card(s, 8.3, 1.95, 2.3, "34", "tỉnh thành")
stat_card(s, 10.75, 1.95, 2.5, "4", "quốc gia")
stat_card(s, 8.3, 3.7, 2.3, "1000+", "review mới/ngày")
stat_card(s, 10.75, 3.7, 2.5, "04/2023", "ra mắt")
txt(s, 8.3, 5.55, 4.45, 1.0, "Điểm cao ≠ an toàn.\nRủi ro nằm ở phần đuôi.", 16, INK, bold=True)
notes(s, "~60s · 14/4/2023 → 34 tỉnh/4 quốc gia → 1000+/ngày. Đọc 2 review thật, nhấn: nằm ngay dưới 4.8★.")

# S3 PAIN POINT & CÁI GIÁ
s = new_slide("HỒI 1 · CÂU CHUYỆN", "Ba lỗ hổng — và cái giá doanh nghiệp phải trả")
card(s, 0.7, 1.95, 3.9, 2.2, "🌊 Quá tải", "Hàng nghìn review/ngày.\nKhông đội ngũ nào đọc xuể\nbằng tay.", fill=AMBER_SOFT, head_color=AMBER)
card(s, 4.75, 1.95, 3.9, 2.2, "⚠️ An toàn bị bỏ sót", "Ca 'say rượu' trông y hệt ca\n'app lag' — người đọc kỹ\nluôn đến sau.", fill=AMBER_SOFT, head_color=AMBER)
card(s, 8.8, 1.95, 3.9, 2.2, "🧭 Sai địa chỉ", "Hộp thư chung: ai cũng thấy,\nkhông ai chịu trách nhiệm.", fill=AMBER_SOFT, head_color=AMBER)
bullets(s, [
    ("Cái giá:", "1 ca an toàn bỏ sót = nguy cơ khủng hoảng truyền thông · 1 lỗi thanh toán không ai nhận = làn sóng rời bỏ · mọi quyết định chất lượng chậm một nhịp so với thực tế."),
], y=4.6, size=17)
txt(s, 0.7, 5.9, 12.0, 0.6, "Đây là kịch bản mặc định của mọi nền tảng có hàng triệu người dùng — không riêng Green SM.", 15, MUTED, italic=True)
notes(s, "~70s · 3 lỗ hổng (nhấn lỗ hổng 2) → quy về chi phí business: truyền thông, churn, quyết định chậm.")

# S4 AI ĐANG ĐAU & CẦN GÌ
s = new_slide("HỒI 1 · CÂU CHUYỆN", "Ai đang đau — và ứng dụng đáp ứng nhu cầu gì")
card(s, 0.7, 2.0, 3.9, 3.3, "Đội vận hành", "Cần biết ca an toàn NGAY\nkhi xuất hiện — không phải\ncuối tuần khi tổng hợp.\n\n→ Cảnh báo realtime,\nSLA 1 giờ cho ca an toàn.")
card(s, 4.75, 2.0, 3.9, 3.3, "Chăm sóc khách hàng", "Cần đúng việc vào đúng\nhàng đợi, kèm ngữ cảnh —\nkhông phải hộp thư hỗn độn.\n\n→ Định tuyến 6 phòng ban,\nAI viết sẵn nguyên nhân gốc.")
card(s, 8.8, 2.0, 3.9, 3.3, "Quản lý", "Cần nhìn toàn cảnh realtime,\ngiao việc có chủ, có hạn,\ncó kiểm đếm.\n\n→ Dashboard + báo cáo AI\n+ vòng đời task khép kín.")
txt(s, 0.7, 5.7, 12.0, 0.6, "Sản phẩm được thiết kế đúng cho ba nhu cầu này — không hơn, không kém.", 16, INK, bold=True)
notes(s, "~60s · 3 personas: vận hành / CSKH / quản lý — mỗi nhóm 1 nhu cầu, 1 câu trả lời của sản phẩm.")

# S5 WHY NOW
s = new_slide("HỒI 1 · CÂU CHUYỆN", "Điều cần tồn tại — và vì sao là bây giờ")
txt(s, 0.7, 2.3, 12.0, 1.6, "Một hệ thống ĐỌC THAY con người,\nPHÂN LOẠI đúng mức khẩn cấp, GIAO ĐÚNG NGƯỜI xử lý.", 30, GREEN, bold=True)
bullets(s, [
    ("Why now:", "AI đọc hiểu tiếng Việt đã đủ tốt và đủ rẻ để đọc 1,5 triệu review — 3 năm trước còn bất khả thi."),
    ("Không phải ý tưởng:", "hệ thống đã chạy thật — phần còn lại của buổi hôm nay là bằng chứng."),
], y=4.4, size=19)
notes(s, "~40s · Chốt Hồi 1 bằng 1 câu giải pháp; bridge sang sản phẩm.")

# ================= HỒI 2 · SẢN PHẨM =================
# S6 GIẢI PHÁP TỔNG QUAN
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Khoang lái điều hành chất lượng theo thời gian thực")
card(s, 0.7, 2.0, 2.95, 3.0, "📥 Tiếp nhận", "Gom phản hồi đa kênh,\nstreaming realtime\n~50 review/giây.")
card(s, 3.8, 2.0, 2.95, 3.0, "🧠 Phân tích", "ML phân loại mọi review;\nLLM đánh giá rủi ro khi\ncần suy luận sâu.")
card(s, 6.9, 2.0, 2.95, 3.0, "🔔 Cảnh báo", "Phát hiện spike, sinh alert,\nđịnh tuyến 6 phòng ban\nkèm SLA.")
card(s, 10.0, 2.0, 2.95, 3.0, "✅ Hành động", "Báo cáo AI, tạo việc,\ngiao việc, khép vòng khi\nhoàn thành.")
txt(s, 0.7, 5.4, 12.0, 0.9, "Từng khối có sản phẩm rõ ràng, kiểm chứng được — máy làm phần nặng, con người giữ quyết định quan trọng.", 16, INK, bold=True)
notes(s, "~60s · 4 khối Tiếp nhận→Phân tích→Cảnh báo→Hành động; câu chốt 'máy làm phần nặng'.")

# S7 DEMO FLOW 5 MÀN HÌNH
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Sản phẩm thật — đi một vòng 5 màn hình")
shots = [("ui-dashboard.png","① Tổng quan"), ("ui-feedback.png","② Phản hồi + nhãn AI"),
         ("ui-ticket.png","③ Ticket 6 phòng ban"), ("ui-report.png","④ Báo cáo AI"), ("ui-tasks.png","⑤ Công việc")]
for i, (f, lab) in enumerate(shots):
    x = 0.45 + (i % 3) * 3.7; y = 1.9 + (i // 3) * 2.65
    pic_fit(s, f"{A}/{f}", x, y, 3.45)
    txt(s, x, y + 2.12, 3.45, 0.35, lab, 11, GREEN, bold=True)
txt(s, 7.85, 4.7, 4.9, 1.6, "Năm màn hình, một mạch duy nhất:\ntừ lời phàn nàn thô đến việc có chủ.", 16, INK, bold=True)
notes(s, "~90s · Đi theo hành trình người dùng qua 5 màn thật; câu chốt 'từ lời phàn nàn thô đến việc có chủ'.")

# S8 HÀNH TRÌNH 7 CHẶNG
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Phòng cấp cứu cho review — hành trình 7 chặng")
chips(s, ["① Che PII", "② ML\nphân loại", "③ Rule\nghi đè", "④ Spike\ndetect", "⑤ Routing\n6 phòng", "⑥ Task\nagent/duyệt", "⑦ Người\nxác nhận"])
bullets(s, [
    ("Hybrid rẻ/đắt:", "ML nhỏ đọc tất cả; LLM chỉ gọi khi cần suy luận sâu."),
    ("An toàn thắng xác suất.", "Rule safety luôn ghi đè được model."),
    ("Config-driven:", "đổi luật = sửa 1 file cấu hình, không sửa code."),
    ("Graceful degradation:", "một khối hỏng, cả hệ vẫn đứng."),
], y=4.35, size=16)
notes(s, "~70s · Ẩn dụ phòng cấp cứu; lướt 7 chip; 4 nguyên tắc vận hành.")

# S9 CHẶNG 1–2
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Chặng 1–2 · Che PII → ML đọc lướt tất cả", ttl_size=28)
bullets(s, [
    ("Che PII trước tiên.", "SĐT, email, biển số bị ẩn trước khi người hay máy đọc — nguyên tắc, không phải tính năng."),
    ("ML nhanh & rẻ, <0,1s/review.", "TF-IDF + LinearSVC: 77,1% chủ đề · 94,3% cảm xúc."),
    ("Hiểu phủ định tiếng Việt.", "“không tệ” không bị đếm là chê."),
    ("Vì rẻ nên đọc được 100%.", "Không ca nào bị bỏ vì thiếu ngân sách."),
], x=0.7, y=1.95, w=6.6, size=16)
pic_fit(s, f"{A}/ui-feedback.png", 7.6, 1.75, 5.3)
notes(s, "~60s · Chặng 1 PII; chặng 2 ML — số 77,1%/94,3%, negation guard; chỉ nhãn AI trên screenshot.")

# S10 CHẶNG 3 — RULE GHI ĐÈ (ví dụ từng bước)
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Chặng 3 · Luật an toàn ghi đè — ví dụ thật từng bước", ttl_size=28)
table(s, 0.7, 1.95, 12.0, 3.4,
      ["Bước", "Kết quả"],
      [["Input (review thật)", "“Tài xế say xỉn gây tai nạn, xe tông người rồi bỏ chạy”"],
       ["Sau che PII", "Giữ nguyên — không có PII"],
       ["ML dự đoán", "Thái độ tài xế · tiêu cực · mức TRUNG BÌNH"],
       ["Rule an toàn GHI ĐÈ", "An toàn · KHẨN CẤP CAO NHẤT · bắt buộc người xem"],
       ["Lý do", "Từ khóa an toàn: “tai nạn”, “bỏ chạy”"]],
      col_w=[3.2, 8.8], font=13)
txt(s, 0.7, 5.8, 12.0, 0.8, "Khi có nghi ngờ về an toàn, máy móc không được phép tự tin hơn sự thận trọng.", 18, GREEN, bold=True)
notes(s, "~70s · Đi chậm từng dòng bảng; câu chốt là câu quan trọng nhất của bài.")

# S11 CHẶNG 4 — SPIKE
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Chặng 4 · Phát hiện ổ dịch — ngưỡng rõ ràng", ttl_size=28)
card(s, 0.7, 2.1, 3.9, 1.9, "1 ca khẩn cấp nhất", "→ Báo động NGAY lập tức.")
card(s, 4.75, 2.1, 3.9, 1.9, "≥ 3 ca mức cao\ncùng khu vực", "→ Cảnh báo VÀNG.", fill=AMBER_SOFT, head_color=AMBER)
card(s, 8.8, 2.1, 3.9, 1.9, "≥ 10 ca tiêu cực", "→ Báo động ĐỎ.", fill=AMBER_SOFT, head_color=AMBER)
bullets(s, [
    ("Gom theo (khu vực, chủ đề, phòng ban).", "Một ca là cá biệt — nhiều ca cùng nơi cùng lúc là Ổ DỊCH."),
    ("Dedup cooldown 60 phút.", "Không báo lại cùng một chuyện — tín hiệu không bị loãng."),
], y=4.5, size=17)
notes(s, "~50s · 3 ngưỡng + cooldown 60'; ẩn dụ ổ dịch.")

# S12 CHẶNG 5 — ROUTER + SLA
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Chặng 5 · Định tuyến đúng chuyên khoa — kèm SLA", ttl_size=28)
table(s, 0.7, 1.95, 7.4, 3.3,
      ["Chủ đề", "Phòng ban", "SLA"],
      [["An toàn (Safety)", "Vận hành", "1 giờ"],
       ["Hành vi tài xế", "Vận hành", "8 giờ"],
       ["Lỗi ứng dụng", "IT", "12 giờ"],
       ["Thanh toán / Giá", "CSKH", "12 giờ"],
       ["P5 (bất kể chủ đề)", "Vận hành ← override", "1 giờ"]],
      col_w=[2.9, 2.7, 1.8], font=13)
card(s, 8.5, 1.95, 4.2, 2.0, "⚙️ Config-driven", "Đổi luật phân khoa =\nsửa 1 file cấu hình.\nKhông sửa code,\nkhông deploy lại.")
card(s, 8.5, 4.15, 4.2, 1.6, "🤖 LLM chỉ viết nội dung", "Tiêu đề, nguyên nhân gốc,\nhành động đề xuất — cho\ncon người đọc nhanh.")
notes(s, "~60s · Bảng phân khoa 6 phòng (bảng hiện 5 dòng tiêu biểu); nhấn 'không viết lại code'.")

# S13 CHẶNG 6 — ALERT→TASK
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Chặng 6 · Từ báo động đến việc thật — hai con đường", ttl_size=28)
card(s, 0.7, 1.95, 5.9, 3.5, "⚙️ Đường tự động — Task agent", "Gom alert cùng (phòng ban, khu vực).\nChưa có task mở → tự lập đầu việc.\nMỗi cụm chỉ 1 task mở —\nkhông trùng, không rơi rớt.")
card(s, 6.85, 1.95, 5.9, 3.5, "✅ Đường có người — duyệt báo cáo", "AI viết báo cáo: nguyên nhân gốc\n+ các đầu việc nháp.\nQuản lý thêm/sửa/xóa → bấm 1 nút:\nduyệt + tạo việc trong cùng một hành động,\nkèm link ngược báo cáo gốc.")
txt(s, 0.7, 5.75, 12.0, 0.7, "MÁY ĐỀ XUẤT. NGƯỜI QUYẾT ĐỊNH. HỆ THỐNG THỰC THI TRỌN VẸN.", 19, GREEN, bold=True, align=PP_ALIGN.CENTER)
notes(s, "~60s · 2 con đường; câu đinh giữa slide.")

# S14 CHẶNG 7 — HITL
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Chặng 7 · Con người trong vòng lặp — và vòng học", ttl_size=28)
bullets(s, [
    ("4 tình huống bắt buộc người xem:", "khẩn cấp nhất / từ khóa an toàn · AI không chắc · câu mơ hồ · có người yêu cầu."),
    ("Người duyệt: đồng ý · sửa nhãn · đánh dấu xong.", "Mỗi lần sửa ghi sổ: ai, trường nào, nhãn cũ → mới, lúc nào."),
    ("Cuốn sổ 2 vai trò:", "hồ sơ kiểm tra lại + dữ liệu cho AI học lại."),
    ("VÒNG HỌC KHÉP KÍN:", "AI gán sai → người sửa → hệ thống ghi nhớ → AI giỏi hơn."),
], x=0.7, y=1.95, w=6.6, size=16)
pic_fit(s, f"{A}/ui-feedback.png", 7.6, 1.75, 5.3)
notes(s, "~70s · Chỉ vào hàng đợi 'Cần duyệt' + nút 'sửa nhãn' trên screenshot.")

# S15 VÒNG ĐỜI TASK
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Vòng đời task — 5 trạng thái, rõ ai làm gì", ttl_size=28)
chips(s, ["Đang mở\nHệ thống tạo\ntừ ticket/báo cáo", "Đã giao\nManager gán người\n+ đặt hạn", "Đang xử lý\nNhân viên bấm\n'Nhận việc'", "Chờ duyệt\nNhân viên nộp\nbáo cáo hoàn thành", "Đã xong\nManager duyệt,\nkhép vòng"], y=2.4, x0=0.6, cw=2.35, gap=0.15)
bullets(s, [
    ("Mỗi lần chuyển trạng thái đều ghi log.", "Phục vụ audit; task quá hạn tự được đánh dấu cho manager."),
], y=4.6, size=17)
notes(s, "~50s · Đọc 5 chip theo thứ tự, nhấn cột 'ai làm'.")

# S16 GIAO VIỆC 2 PHÍA
s = new_slide("HỒI 2 · SẢN PHẨM THỰC TẾ", "Vận hành giao việc — hai phía màn hình", ttl_size=28)
pic_fit(s, f"{A}/ui-tasks.png", 0.6, 1.85, 5.5)
txt(s, 0.6, 5.45, 5.9, 1.0, "Manager: task gắn sẵn phòng ban → 'Gán task'\n→ AI gợi ý người nhận cân bằng tải → đặt hạn → Gán.", 12, INK)
pic_fit(s, f"{A}/ui-staff.png", 6.85, 1.85, 5.5)
txt(s, 6.85, 5.45, 5.9, 1.0, "Nhân viên: 'Việc của tôi' chỉ hiện việc của mình →\n'Nhận việc' → nộp báo cáo (Hoàn thành/Một phần/Bị chặn).", 12, INK)
txt(s, 0.6, 6.55, 12.1, 0.5, "Không ai phải hỏi “việc này của ai” — hệ thống trả lời bằng thiết kế.", 15, GREEN, bold=True)
notes(s, "~70s · 2 screenshot 2 phía; câu chốt cuối slide.")

# ================= HỒI 3 · BẰNG CHỨNG =================
# S17 SỐ ĐO
s = new_slide("HỒI 3 · BẰNG CHỨNG & KỸ THUẬT", "Chạy thật, đo thật — không phải số trình diễn")
stats = [("77,1%", "đúng chủ đề (932 mẫu test)"), ("94,3%", "đúng cảm xúc"), ("97%", "không bỏ sót tiêu cực"),
         ("97 giây", "xử lý 4.837 review"), ("~50/giây", "tốc độ streaming"), ("6/6", "kịch bản kiểm thử đạt")]
for i, (big, small) in enumerate(stats):
    stat_card(s, 0.7 + (i % 3) * 4.2, 2.1 + (i // 3) * 2.0, 3.9, big, small)
txt(s, 0.7, 6.3, 12.0, 0.5, "Mọi con số tái lập được bằng bộ kiểm thử trong mã nguồn.", 15, INK, bold=True)
notes(s, "~60s · Câu đời thường trước, phần trăm sau; 6 stat card.")

# S18 BẢNG TEST CASE
s = new_slide("HỒI 3 · BẰNG CHỨNG & KỸ THUẬT", "6/6 kịch bản kiểm thử — mỗi dòng một lời hứa của thiết kế", ttl_size=26)
table(s, 0.7, 1.9, 12.0, 4.4,
      ["Test case", "Input (rút gọn)", "Output thật", "Đạt"],
      [["TC-01 Safety override", "“Tài xế say xỉn, xe tông người bỏ chạy”", "An toàn · P5 · bắt buộc duyệt", "✅"],
       ["TC-02 Lỗi app", "“App bị lỗi không đặt được xe”", "Lỗi ứng dụng · P3", "✅"],
       ["TC-03 Gian lận thanh toán", "“Tài xế thu tiền gấp đôi”", "Thanh toán · P3", "✅"],
       ["TC-04 Positive guard", "“Tài xế thân thiện, lái an toàn”", "Tài xế · tích cực · P1", "✅"],
       ["TC-05 Câu gián tiếp", "“Gọi mãi không có tài xế nhận”", "Khác · tiêu cực · P3", "✅"],
       ["TC-06 Che PII", "SĐT 0912… · email … · biển 30A…", "[SĐT_ĐÃ_ẨN] [EMAIL_ĐÃ_ẨN] …", "✅"]],
      col_w=[2.6, 4.6, 3.6, 1.2], font=12)
notes(s, "~50s · Không đọc từng ô — kể 6 lời hứa: an toàn, đúng phòng, bắt gian lận, không báo oan, hiểu câu gián tiếp, che PII.")

# S19 KIẾN TRÚC
s = new_slide("HỒI 3 · BẰNG CHỨNG & KỸ THUẬT", "Kiến trúc — 6 dịch vụ độc lập, LLM có dự phòng 4 lớp", ttl_size=26)
chips(s, ["🖥️ Frontend\nReact 19\n1 service", "⚙️ Backend API\nFastAPI\n1 service", "🔧 Ingest\nworker", "🔧 ML\nworker", "🔧 Ticket\nworker", "🔧 Agent\nworker", "🗄️ Supabase\nPostgreSQL\ncloud DB"], y=2.3, x0=0.45, cw=1.72, gap=0.1)
card(s, 0.7, 4.3, 12.0, 1.7, "⚡ LLM router duy nhất — failover 4 lớp", "OpenRouter → OpenAI → Gemini → Groq (→ mock khi demo offline).\nMột nhà cung cấp sập, hệ thống tự chuyển — buổi demo không thể vỡ vì bên thứ ba.")
notes(s, "~60s · 6 service + DB; nhấn failover: 'demo không thể vỡ'.")

# S20 TECH STACK
s = new_slide("HỒI 3 · BẰNG CHỨNG & KỸ THUẬT", "Tech stack — nhàm chán mà chắc", ttl_size=28)
card(s, 0.7, 2.0, 3.9, 3.2, "Frontend", "React 19\nTanStack Start/Router/Query\nVite 7 · Tailwind CSS 4\nshadcn/ui · Recharts")
card(s, 4.75, 2.0, 3.9, 3.2, "Backend / AI", "FastAPI · Python\nscikit-learn:\nTF-IDF + LinearSVC\nAPScheduler")
card(s, 8.8, 2.0, 3.9, 3.2, "Hạ tầng & LLM", "Supabase (PostgreSQL, Auth)\nRailway — 6 Docker service\nOpenRouter/OpenAI/\nGemini/Groq failover")
txt(s, 0.7, 5.5, 12.0, 0.6, "Không công nghệ nào ở đây cần một chuyên gia hiếm để vận hành.", 15, MUTED, italic=True)
notes(s, "~40s · Lướt nhanh; câu chốt 'nhàm chán mà chắc'.")

# ================= HỒI 4 · GIÁ TRỊ & TƯƠNG LAI =================
# S21 TƯ DUY SẢN PHẨM
s = new_slide("HỒI 4 · GIÁ TRỊ & TƯƠNG LAI", "Tư duy sản phẩm — 5 nguyên tắc định hình mọi quyết định", ttl_size=26)
rows5 = [("01 · An toàn không nhường xác suất", "Rule safety ép mức cao nhất dù model nghĩ khác."),
         ("02 · Con người trong vòng lặp", "Máy đề xuất, người quyết định & chịu trách nhiệm."),
         ("03 · Đúng người đúng dữ liệu", "Quyền thao tác kiểm ở backend theo vai trò & phòng ban."),
         ("04 · Khép kín phản hồi → hành động", "Phân tích không tạo ra hành động là phân tích trang trí."),
         ("05 · Realtime trước tiên", "Thấy vấn đề khi nó đang xảy ra, không phải hôm sau.")]
for i, (h, b) in enumerate(rows5):
    y = 1.95 + i * 0.95
    rect(s, 0.7, y, 12.0, 0.82, GREEN_SOFT if i % 2 == 0 else ROW_ALT, line_color=GREEN, rounded=True)
    txt(s, 0.95, y + 0.08, 5.4, 0.65, h, 15, GREEN, bold=True)
    txt(s, 6.5, y + 0.08, 6.0, 0.65, b, 13, INK)
notes(s, "~60s · 5 nguyên tắc; nhấn 04: 'phân tích trang trí'.")

# S22 ỨNG DỤNG THỰC TẾ — PAIN → NĂNG LỰC → KẾT QUẢ
s = new_slide("HỒI 4 · GIÁ TRỊ & TƯƠNG LAI", "Ứng dụng thực tế — pain point nào được giải quyết", ttl_size=26)
table(s, 0.7, 1.95, 12.0, 3.6,
      ["Pain point", "Năng lực hệ thống", "Kết quả nhận được"],
      [["Quá tải review", "ML đọc 100%, người chỉ duyệt ca máy không chắc", "Hàng đợi duyệt thay vì biển dữ liệu"],
       ["An toàn bị bỏ sót", "Rule ghi đè + alert SLA 1 giờ cho ca an toàn", "Phát hiện khi đang xảy ra, không phải khi lên báo"],
       ["Sai địa chỉ", "Định tuyến 6 phòng + giao việc từng người, có hạn", "Luôn có một cái tên chịu trách nhiệm"]],
      col_w=[2.9, 4.9, 4.2], font=13)
txt(s, 0.7, 5.9, 12.0, 0.7, "Ba nhu cầu của vận hành, CSKH, quản lý — trả lời bằng thiết kế, không phải bằng lời hứa.", 16, GREEN, bold=True)
notes(s, "~70s · Móc lại S3–S4: từng pain point → năng lực → kết quả.")

# S23 HẠN CHẾ & ROADMAP
s = new_slide("HỒI 4 · GIÁ TRỊ & TƯƠNG LAI", "Chúng tôi biết mình đang ở đâu", ttl_size=28)
bullets(s, [
    ("P0 · Khóa sâu tầng dữ liệu.", "Cửa chính (tầng ứng dụng) đã gate chặt; lớp khóa tại kho dữ liệu là bước kế — biết chính xác việc cần làm."),
    ("P1 · Nối dữ liệu sống.", "Đang chạy tập cố định ~5.000 review thật; kế tiếp: ống dữ liệu hằng ngày, giữ nguyên phần xử lý."),
    ("P1 · Báo động mới chỉ trên web.", "Người trực không mở máy 2 giờ sáng vẫn lỡ. Kế tiếp: đẩy thẳng vào điện thoại."),
], y=2.1, size=19)
txt(s, 0.7, 5.6, 12.0, 0.6, "Minh bạch về giới hạn = dấu hiệu của sự nghiêm túc, không phải điểm yếu.", 15, MUTED, italic=True)
notes(s, "~70s · Giọng tự tin. Ý 3 móc lại mở đầu — nhấn chậm.")

# S24 VÒNG LẶP ĐÃ KHÉP
s = new_slide("HỒI 4 · GIÁ TRỊ & TƯƠNG LAI", "Vòng lặp đã khép", ttl_size=32)
chips(s, ["📥 Review\nvào", "🤖 AI\nphân loại", "🔔 Báo động\nnổi đúng lúc", "✅ Quản lý\nduyệt", "💼 Nhân viên\nxử lý", "📋 Hoàn thành\n↺ quay lại"], y=2.6, x0=0.6, cw=1.95, gap=0.12)
txt(s, 0.7, 4.8, 12.0, 1.0, "Tự động, có kiểm soát — không ai phải thức trắng đêm dò từng dòng review.", 20, INK, bold=True, align=PP_ALIGN.CENTER)
notes(s, "~40s · Callback 2 giờ sáng; đọc vòng 6 chip.")

# S25 VISION & KẾT (full green)
s = prs.slides.add_slide(BLANK); PAGE[0] += 1
rect(s, 0, 0, 13.333, 7.5, GREEN)
txt(s, 0.9, 1.5, 11.5, 1.4, "Một lời phàn nàn không được đọc\nlà một sự cố đang chờ xảy ra lần nữa.", 34, WHITE, bold=True)
txt(s, 0.9, 3.7, 11.5, 1.3, "Tầm nhìn 3–5 năm: tầng lắng nghe tiêu chuẩn cho mọi nền tảng tiêu dùng Việt Nam —\nnơi mọi lời phàn nàn được đọc, được định tuyến, và trở thành hành động.", 18, WHITE)
txt(s, 0.9, 5.7, 11.5, 1.0, "Với chúng tôi — nó sẽ được đọc.   ·   Team 116 · AI20K Cohort 2 · Q&A", 20, WHITE, bold=True)
notes(s, "~60s · Vision → câu kết → cảm ơn + mời Q&A.")

out = "docs/project/slides/GreenSM-Master-Deck.pptx"
prs.save(out); print("saved", out, "slides:", len(prs.slides._sldIdLst))
